import json
import subprocess
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

from docx.oxml import OxmlElement
from docx.oxml.ns import qn


BASE_DIR = Path(__file__).resolve().parent.parent

PROFILE_FILE = BASE_DIR / "profile" / "master_profile.json"
ADAPTED_FILE = BASE_DIR / "output" / "adapted_profile.json"
JOB_FILE = BASE_DIR / "output" / "job_profile.json"
MATCH_FILE = BASE_DIR / "output" / "match_result.json"

OUTPUT_FILE = BASE_DIR / "output" / "CV_Gonzalo_Vega_Ramos.docx"
PDF_OUTPUT_FILE = BASE_DIR / "output" / "CV_Gonzalo_Vega_Ramos.pdf"


def load_json(path: Path) -> dict:
    with open(path, "r", encoding="utf-8") as file:
        return json.load(file)


def add_bottom_border(paragraph):
    p = paragraph._p
    p_pr = p.get_or_add_pPr()

    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")

    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "808080")

    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def add_section_title(document, title):
    paragraph = document.add_paragraph()

    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(4)

    run = paragraph.add_run(title.upper())
    run.bold = True
    run.font.size = Pt(11)

    add_bottom_border(paragraph)

    return paragraph


def add_bullet(document, text):
    paragraph = document.add_paragraph(style="List Bullet")

    paragraph.paragraph_format.space_after = Pt(2)

    run = paragraph.add_run(str(text))
    run.font.size = Pt(9)

    return paragraph


def add_experience(document, experience):
    paragraph = document.add_paragraph()

    paragraph.paragraph_format.space_after = Pt(2)

    company = paragraph.add_run(
        f'{experience.get("company", "")} — '
    )
    company.bold = True
    company.font.size = Pt(10)

    role = paragraph.add_run(
        experience.get("role", "")
    )
    role.bold = True
    role.font.size = Pt(10)

    duration = experience.get("duration")

    if duration:
        paragraph.add_run(
            f' | {duration}'
        ).font.size = Pt(9)

    technologies = experience.get("technologies", [])

    if technologies:
        tech_paragraph = document.add_paragraph()

        tech_paragraph.paragraph_format.space_after = Pt(2)

        run = tech_paragraph.add_run(
            "Tecnologías: "
        )
        run.bold = True
        run.font.size = Pt(8.5)

        tech_run = tech_paragraph.add_run(
            ", ".join(technologies)
        )
        tech_run.font.size = Pt(8.5)

    for achievement in experience.get("achievements", []):
        add_bullet(document, achievement)


def add_project(document, project):
    paragraph = document.add_paragraph()

    paragraph.paragraph_format.space_after = Pt(2)

    name = paragraph.add_run(
        project.get("name", "")
    )
    name.bold = True
    name.font.size = Pt(10)

    project_type = project.get("type")

    if project_type:
        paragraph.add_run(
            f" — {project_type}"
        ).font.size = Pt(9)

    description = project.get("description")

    if description:
        paragraph = document.add_paragraph(description)
        paragraph.paragraph_format.space_after = Pt(2)

        for run in paragraph.runs:
            run.font.size = Pt(9)

    technologies = project.get("technologies", [])

    if technologies:
        paragraph = document.add_paragraph()

        run = paragraph.add_run("Tecnologías: ")
        run.bold = True
        run.font.size = Pt(8.5)

        run = paragraph.add_run(
            ", ".join(technologies)
        )
        run.font.size = Pt(8.5)

    for achievement in project.get("achievements", []):
        add_bullet(document, achievement)


def add_skills(document, skills):
    """
    Renderiza las skills agrupadas por categoría.
    """

    if not skills:
        return

    add_section_title(
        document,
        "Competencias técnicas"
    )

    category_names = {
        "backend": "Backend",
        "programming": "Programación",
        "ai_ml": "IA / Machine Learning",
        "cloud": "Cloud",
        "data": "Datos",
        "tools": "Herramientas",
        "frontend": "Frontend"
    }

    for category, values in skills.items():

        if not values:
            continue

        label = category_names.get(
            category,
            category
        )

        paragraph = document.add_paragraph()

        paragraph.paragraph_format.space_after = Pt(2)

        run = paragraph.add_run(
            f"{label}: "
        )
        run.bold = True
        run.font.size = Pt(9)

        run = paragraph.add_run(
            ", ".join(values)
        )
        run.font.size = Pt(9)

def add_ats_keywords(document, keywords):
    """
    Renderiza keywords adicionales orientadas a ATS.
    Estas keywords no forman parte de las skills originales
    del candidato.
    """

    if not keywords:
        return

    add_section_title(
        document,
        "Competencias adicionales"
    )

    paragraph = document.add_paragraph()

    paragraph.paragraph_format.space_after = Pt(2)

    run = paragraph.add_run(
        ", ".join(keywords)
    )

    run.font.size = Pt(9)


def add_education(document, education):
    """
    Renderiza formación tanto si llega como texto
    como si llega como objeto.
    """

    if not education:
        return

    add_section_title(
        document,
        "Formación"
    )

    for item in education:

        # Formato simple:
        # "Grado Superior en Desarrollo..."
        if isinstance(item, str):
            add_bullet(
                document,
                item
            )
            continue

        # Formato estructurado:
        # {
        #     "title": "...",
        #     "institution": "..."
        # }
        if isinstance(item, dict):

            title = item.get(
                "title",
                ""
            )

            institution = item.get(
                "institution"
            )

            text = title

            if institution:
                text += f" — {institution}"

            add_bullet(
                document,
                text
            )


def add_languages(document, languages):
    """
    Renderiza idiomas desde el formato del CVAdapter:

    [
        {
            "name": "Español",
            "level": "native"
        }
    ]
    """

    if not languages:
        return

    add_section_title(
        document,
        "Idiomas"
    )

    for language in languages:

        if isinstance(language, dict):

            name = language.get(
                "name",
                ""
            )

            level = language.get(
                "level",
                ""
            )

            add_bullet(
                document,
                f"{name}: {level}"
            )

        elif isinstance(language, str):

            add_bullet(
                document,
                language
            )


def generate_cv(
    profile,
    adapted,
    job,
    match,
    output_dir
):

    candidate = profile["candidate"]

    document = Document()

    # --------------------------------------------------
    # Página
    # --------------------------------------------------

    section = document.sections[0]

    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    # --------------------------------------------------
    # Fuente global
    # --------------------------------------------------

    styles = document.styles

    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9)

    # --------------------------------------------------
    # CABECERA
    # --------------------------------------------------

    paragraph = document.add_paragraph()

    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(2)

    run = paragraph.add_run(
        candidate["name"]
    )

    run.bold = True
    run.font.size = Pt(20)

    # --------------------------------------------------
    # PUESTO OBJETIVO
    # --------------------------------------------------

    paragraph = document.add_paragraph()

    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(2)

    run = paragraph.add_run(
        adapted.get(
            "target_position",
            job.get("title", "")
        )
    )

    run.bold = True
    run.font.size = Pt(11)

    # --------------------------------------------------
    # CONTACTO
    # --------------------------------------------------

    paragraph = document.add_paragraph()

    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(8)

    contact = [
        candidate.get("location"),
        candidate.get("email"),
        candidate.get("phone"),
        candidate.get("linkedin")
    ]

    contact = [
        item for item in contact
        if item
    ]

    run = paragraph.add_run(
        " | ".join(contact)
    )

    run.font.size = Pt(8.5)

    # --------------------------------------------------
    # PERFIL
    # --------------------------------------------------

    summary = adapted.get("summary")

    if summary:

        add_section_title(
            document,
            "Perfil profesional"
        )

        paragraph = document.add_paragraph(
            summary
        )

        paragraph.paragraph_format.space_after = Pt(5)

        for run in paragraph.runs:
            run.font.size = Pt(9)

    # --------------------------------------------------
    # SKILLS
    # --------------------------------------------------

    skills = adapted.get(
        "skills",
        {}
    )

    add_skills(
        document,
        skills
    )

    # --------------------------------------------------
    # ATS KEYWORDS
    # --------------------------------------------------

    ats_keywords = adapted.get(
        "ats_keywords",
        []
    )

    add_ats_keywords(
        document,
        ats_keywords
    )

    # --------------------------------------------------
    # EXPERIENCIA
    # --------------------------------------------------

    experience = adapted.get(
        "experience",
        []
    )

    if experience:

        add_section_title(
            document,
            "Experiencia profesional"
        )

        for item in experience:

            add_experience(
                document,
                item
            )

    # --------------------------------------------------
    # PROYECTOS
    # --------------------------------------------------

    projects = adapted.get(
        "projects",
        []
    )

    if projects:

        add_section_title(
            document,
            "Proyectos"
        )

        for project in projects:

            add_project(
                document,
                project
            )

    # --------------------------------------------------
    # FORMACIÓN
    # --------------------------------------------------

    education = adapted.get(
        "education",
        []
    )

    add_education(
        document,
        education
    )

    # --------------------------------------------------
    # IDIOMAS
    # --------------------------------------------------

    languages = adapted.get(
        "languages",
        []
    )

    add_languages(
        document,
        languages
    )

    # --------------------------------------------------
    # SALIDA
    # --------------------------------------------------

    output_dir = Path(output_dir)

    output_dir.mkdir(
        parents=True,
        exist_ok=True
    )

    output_file = (
        output_dir /
        "CV_Gonzalo_Vega_Ramos.docx"
    )

    pdf_output_file = (
        output_dir /
        "CV_Gonzalo_Vega_Ramos.pdf"
    )

    # --------------------------------------------------
    # Guardar DOCX
    # --------------------------------------------------

    document.save(
        output_file
    )

    # --------------------------------------------------
    # Convertir a PDF
    # --------------------------------------------------

    convert_to_pdf(
        output_file,
        output_dir
    )

    print(
        f"CV DOCX generado: {output_file}"
    )

    print(
        f"CV PDF generado: {pdf_output_file}"
    )

    print(
        f'Match de la oferta: '
        f'{match["summary"]["match_percentage"]}%'
    )

    return {
        "docx": output_file,
        "pdf": pdf_output_file
    }


def convert_to_pdf(
    docx_file,
    output_dir
):

    soffice = (
        r"C:\Program Files\LibreOffice\program\soffice.exe"
    )

    subprocess.run(
        [
            soffice,
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(output_dir),
            str(docx_file)
        ],
        check=True
    )


if __name__ == "__main__":
    generate_cv()